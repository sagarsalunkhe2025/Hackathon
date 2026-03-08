"""Generate Role Assistant Hackathon PowerPoint presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colors
PRIMARY = RGBColor(99, 102, 241)    # Indigo #6366f1
DARK = RGBColor(30, 30, 46)         # Dark background
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(240, 240, 245)
TEXT_DARK = RGBColor(40, 40, 60)
TEXT_MED = RGBColor(100, 100, 120)
SUCCESS = RGBColor(16, 185, 129)    # Green
WARNING = RGBColor(245, 158, 11)    # Amber
DANGER = RGBColor(239, 68, 68)      # Red
INFO_BLUE = RGBColor(59, 130, 246)  # Blue
ACCENT = RGBColor(139, 92, 246)     # Purple accent

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


def add_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape_bg(slide, left, top, width, height, color, alpha=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=TEXT_DARK, bold=False, align=PP_ALIGN.LEFT, font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_paragraph(text_frame, text, font_size=16, color=TEXT_DARK, bold=False,
                  align=PP_ALIGN.LEFT, space_before=Pt(4), space_after=Pt(4), font_name="Segoe UI"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    p.space_before = space_before
    p.space_after = space_after
    return p


def add_bullet(text_frame, text, font_size=15, color=TEXT_DARK, bold=False, level=0):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = "Segoe UI"
    p.level = level
    p.space_before = Pt(4)
    p.space_after = Pt(4)
    return p


def add_table(slide, left, top, width, height, rows, cols, data, col_widths=None):
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w

    # Header row styling
    for j in range(cols):
        cell = table.cell(0, j)
        cell.text = data[0][j]
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.font.name = "Segoe UI"
            p.alignment = PP_ALIGN.LEFT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    # Data rows
    for i in range(1, rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = data[i][j]
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 1 else LIGHT_GRAY
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.color.rgb = TEXT_DARK
                p.font.name = "Segoe UI"
                p.alignment = PP_ALIGN.LEFT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    return table_shape


def slide_number_footer(slide, num, total):
    add_text_box(slide, Inches(12.2), Inches(7.05), Inches(1), Inches(0.4),
                 f"{num}/{total}", font_size=10, color=TEXT_MED, align=PP_ALIGN.RIGHT)


# ============================================================
# SLIDE 1: Title Slide
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide, DARK)

# Accent bar at top
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.08), PRIMARY)

# Shield icon placeholder
add_text_box(slide, Inches(4.5), Inches(1.5), Inches(4.5), Inches(1.2),
             "\U0001f6e1\ufe0f", font_size=60, color=WHITE, align=PP_ALIGN.CENTER)

# Title
add_text_box(slide, Inches(1.5), Inches(2.6), Inches(10.3), Inches(1.2),
             "Role Assistant", font_size=52, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# Subtitle
add_text_box(slide, Inches(2), Inches(3.8), Inches(9.3), Inches(0.8),
             "AI-Powered SAP Role Governance & Risk Analysis",
             font_size=24, color=RGBColor(180, 180, 220), align=PP_ALIGN.CENTER)

# Tagline
add_text_box(slide, Inches(3), Inches(5.0), Inches(7.3), Inches(0.6),
             "Ask your data, don\u2019t query it.",
             font_size=18, color=PRIMARY, align=PP_ALIGN.CENTER, font_name="Segoe UI Semibold")

# Team
add_text_box(slide, Inches(4), Inches(5.8), Inches(5.3), Inches(0.6),
             "Hackathon Presentation \u2022 2026",
             font_size=16, color=TEXT_MED, align=PP_ALIGN.CENTER)

# Bottom accent bar
add_shape_bg(slide, Inches(0), Inches(7.42), W, Inches(0.08), PRIMARY)


# ============================================================
# SLIDE 2: Agenda
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), PRIMARY)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.7),
             "Agenda", font_size=36, color=DARK, bold=True)

agenda_items = [
    ("01", "Problem Statement", "SAP role governance challenges"),
    ("02", "What We Aim to Achieve", "Vision and success metrics"),
    ("03", "Solution Overview", "Role Assistant capabilities"),
    ("04", "Architecture & Flow", "How the system works end-to-end"),
    ("05", "Key Features", "NL chat, analyses, risk scoring"),
    ("06", "Business Impact", "ROI and compliance alignment"),
    ("07", "Future Roadmap", "What\u2019s next for Role Assistant"),
]

for idx, (num, title, desc) in enumerate(agenda_items):
    y = Inches(1.5) + Inches(idx * 0.78)
    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), y, Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = PRIMARY
    circle.line.fill.background()
    tf = circle.text_frame
    tf.paragraphs[0].text = num
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.name = "Segoe UI"
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    add_text_box(slide, Inches(2.0), y - Inches(0.02), Inches(5), Inches(0.35),
                 title, font_size=20, color=DARK, bold=True)
    add_text_box(slide, Inches(2.0), y + Inches(0.3), Inches(6), Inches(0.3),
                 desc, font_size=14, color=TEXT_MED)

slide_number_footer(slide, 2, 10)


# ============================================================
# SLIDE 3: The Problem
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), DANGER)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(8), Inches(0.7),
             "The Problem: SAP Role Governance is Broken", font_size=32, color=DARK, bold=True)

add_text_box(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.5),
             "Enterprise SAP environments accumulate thousands of roles and millions of user-role assignments over time.",
             font_size=16, color=TEXT_MED)

# Problem cards - 2 rows of 3
problems = [
    ("\U0001f501", "Duplicate Roles", "Identical permissions, different names \u2014 wasted effort & audit confusion", DANGER),
    ("\U0001f6ab", "Unused Roles", "Never assigned to any user \u2014 bloated landscape, increased attack surface", WARNING),
    ("\u26a0\ufe0f", "Over-Privileged Users", "20+ roles or ALL transaction codes \u2014 SOD violations & audit failures", DANGER),
    ("\u23f0", "Expired Assignments", "Past expiry date but still active \u2014 regulatory non-compliance", WARNING),
    ("\U0001f517", "Orphaned Assignments", "Mappings to deleted roles \u2014 data integrity breakdown", WARNING),
    ("\U0001f512", "No Self-Service", "Business users depend on IT \u2014 weeks of delay for security insights", INFO_BLUE),
]

for idx, (icon, title, desc, color) in enumerate(problems):
    col = idx % 3
    row = idx // 3
    x = Inches(0.8) + Inches(col * 4.05)
    y = Inches(1.8) + Inches(row * 2.2)

    # Card background
    card = add_shape_bg(slide, x, y, Inches(3.8), Inches(1.9), LIGHT_GRAY)
    card.shadow.inherit = False

    # Left color bar
    add_shape_bg(slide, x, y, Inches(0.08), Inches(1.9), color)

    # Icon + Title
    add_text_box(slide, x + Inches(0.25), y + Inches(0.15), Inches(3.4), Inches(0.45),
                 f"{icon}  {title}", font_size=18, color=DARK, bold=True)

    # Description
    add_text_box(slide, x + Inches(0.25), y + Inches(0.65), Inches(3.3), Inches(1.1),
                 desc, font_size=13, color=TEXT_MED)

# Quote at bottom
quote_bg = add_shape_bg(slide, Inches(0.8), Inches(6.3), Inches(11.7), Inches(0.9), DARK)
add_text_box(slide, Inches(1.2), Inches(6.4), Inches(11), Inches(0.7),
             "\u201cOrganizations spend $2.5M+ annually on manual SAP role reviews. "
             "Security teams lack real-time visibility, and business users cannot self-serve.\u201d",
             font_size=14, color=RGBColor(200, 200, 220), align=PP_ALIGN.CENTER)

slide_number_footer(slide, 3, 10)


# ============================================================
# SLIDE 4: What We Aim to Achieve
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), PRIMARY)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "What We Aim to Achieve", font_size=36, color=DARK, bold=True)

add_text_box(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(0.5),
             "Vision: Democratize SAP Role Intelligence", font_size=20, color=PRIMARY, bold=True)

# Goals
goals = [
    ("\U0001f4ac", "Natural Language Access",
     "Let business users ask role governance questions in plain English \u2014 no SQL knowledge required"),
    ("\u26a1", "Instant Risk Detection",
     "Surface over-privileged users, duplicates, and expired assignments in seconds \u2014 not weeks"),
    ("\U0001f9e0", "AI-Powered Analysis",
     "Use LLMs to generate SQL queries, analyze results, and provide actionable insights with severity ratings"),
    ("\U0001f6e1\ufe0f", "Proactive Governance",
     "Shift from reactive audit responses to proactive, continuous role monitoring"),
    ("\U0001f4ca", "Quantitative Risk Scoring",
     "Assign risk scores to users based on role count, privileges, and expired assignments"),
]

for idx, (icon, title, desc) in enumerate(goals):
    y = Inches(1.8) + Inches(idx * 0.95)
    add_shape_bg(slide, Inches(0.8), y, Inches(0.06), Inches(0.7), PRIMARY)
    add_text_box(slide, Inches(1.1), y, Inches(10), Inches(0.4),
                 f"{icon}  {title}", font_size=18, color=DARK, bold=True)
    add_text_box(slide, Inches(1.4), y + Inches(0.35), Inches(10), Inches(0.4),
                 desc, font_size=14, color=TEXT_MED)

# Success Metrics box
metrics_bg = add_shape_bg(slide, Inches(7.5), Inches(1.8), Inches(5.2), Inches(4.7), RGBColor(238, 242, 255))
add_text_box(slide, Inches(7.8), Inches(1.9), Inches(4.5), Inches(0.5),
             "Success Metrics", font_size=20, color=PRIMARY, bold=True)

metrics = [
    ("\u23f1\ufe0f  Time-to-Insight", "Weeks \u2192 Seconds"),
    ("\U0001f50d  Coverage", "7 pre-built + unlimited NL queries"),
    ("\U0001f3af  Accuracy", "AI SQL validated & sandboxed"),
    ("\U0001f465  Accessibility", "Zero SQL knowledge needed"),
]

for idx, (label, value) in enumerate(metrics):
    y = Inches(2.5) + Inches(idx * 0.85)
    add_text_box(slide, Inches(7.8), y, Inches(4.8), Inches(0.35),
                 label, font_size=15, color=DARK, bold=True)
    add_text_box(slide, Inches(7.8), y + Inches(0.35), Inches(4.8), Inches(0.35),
                 value, font_size=14, color=TEXT_MED)

slide_number_footer(slide, 4, 10)


# ============================================================
# SLIDE 5: Solution Overview
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), PRIMARY)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Solution Overview", font_size=36, color=DARK, bold=True)

add_text_box(slide, Inches(0.8), Inches(1.2), Inches(11), Inches(0.8),
             "Role Assistant is an intelligent conversational interface that sits on top of your SAP role data. "
             "Users type questions in natural language, and the system translates, executes, analyzes, and presents.",
             font_size=16, color=TEXT_MED)

# 4-step flow
steps = [
    ("1", "Translate", "Question \u2192 safe T-SQL\nvia GPT-4o", PRIMARY),
    ("2", "Execute", "Run query against\nSAP role database", SUCCESS),
    ("3", "Analyze", "AI generates\nhuman-readable summary", ACCENT),
    ("4", "Present", "Insights + data table\n+ generated SQL", INFO_BLUE),
]

for idx, (num, title, desc, color) in enumerate(steps):
    x = Inches(0.8) + Inches(idx * 3.15)
    y = Inches(2.3)

    card = add_shape_bg(slide, x, y, Inches(2.8), Inches(1.8), color)

    add_text_box(slide, x + Inches(0.2), y + Inches(0.15), Inches(2.4), Inches(0.5),
                 f"Step {num}", font_size=14, color=RGBColor(220, 220, 255), bold=False)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.5), Inches(2.4), Inches(0.5),
                 title, font_size=22, color=WHITE, bold=True)
    add_text_box(slide, x + Inches(0.2), y + Inches(1.0), Inches(2.4), Inches(0.7),
                 desc, font_size=13, color=RGBColor(230, 230, 255))

    # Arrow between steps
    if idx < 3:
        add_text_box(slide, x + Inches(2.8), y + Inches(0.6), Inches(0.35), Inches(0.5),
                     "\u25b6", font_size=20, color=TEXT_MED, align=PP_ALIGN.CENTER)

# Two modes section
add_text_box(slide, Inches(0.8), Inches(4.5), Inches(5), Inches(0.5),
             "Two Modes of Operation", font_size=22, color=DARK, bold=True)

# Chat Mode card
chat_bg = add_shape_bg(slide, Inches(0.8), Inches(5.1), Inches(5.8), Inches(2.1), RGBColor(238, 242, 255))
add_shape_bg(slide, Inches(0.8), Inches(5.1), Inches(0.08), Inches(2.1), PRIMARY)
add_text_box(slide, Inches(1.1), Inches(5.2), Inches(5.3), Inches(0.4),
             "\U0001f4ac  Chat Mode", font_size=20, color=PRIMARY, bold=True)
txBox = add_text_box(slide, Inches(1.1), Inches(5.6), Inches(5.3), Inches(1.4),
                     "", font_size=14, color=TEXT_MED)
tf = txBox.text_frame
tf.paragraphs[0].text = "Ask anything in English:"
tf.paragraphs[0].font.size = Pt(14)
tf.paragraphs[0].font.color.rgb = TEXT_MED
tf.paragraphs[0].font.name = "Segoe UI"
for q in ['"Show users with >10 roles"', '"Find roles in Finance area"', '"Who has full access?"']:
    add_bullet(tf, q, font_size=13, color=TEXT_DARK)

# Quick Analysis card
qa_bg = add_shape_bg(slide, Inches(6.9), Inches(5.1), Inches(5.8), Inches(2.1), RGBColor(236, 253, 245))
add_shape_bg(slide, Inches(6.9), Inches(5.1), Inches(0.08), Inches(2.1), SUCCESS)
add_text_box(slide, Inches(7.2), Inches(5.2), Inches(5.3), Inches(0.4),
             "\U0001f4ca  Quick Analysis Mode", font_size=20, color=SUCCESS, bold=True)
txBox2 = add_text_box(slide, Inches(7.2), Inches(5.6), Inches(5.3), Inches(1.4),
                      "", font_size=14, color=TEXT_MED)
tf2 = txBox2.text_frame
tf2.paragraphs[0].text = "One-click pre-built analyses:"
tf2.paragraphs[0].font.size = Pt(14)
tf2.paragraphs[0].font.color.rgb = TEXT_MED
tf2.paragraphs[0].font.name = "Segoe UI"
for a in ["Duplicate & Unused Roles", "Over-Privileged Users", "Expired & Orphaned Assignments",
          "Role Distribution & User Risk Scoring"]:
    add_bullet(tf2, a, font_size=13, color=TEXT_DARK)

slide_number_footer(slide, 5, 10)


# ============================================================
# SLIDE 6: Architecture
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), PRIMARY)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Architecture & Flow", font_size=36, color=DARK, bold=True)

# Architecture boxes
components = [
    (Inches(0.6), Inches(1.6), Inches(3.2), Inches(2.6), "FRONTEND", "React 18 + TypeScript\nVite\n\n\u2022 Chat UI\n\u2022 Data Tables\n\u2022 Quick Analysis Sidebar\n\u2022 Category Filters", INFO_BLUE),
    (Inches(4.5), Inches(1.6), Inches(3.8), Inches(2.6), "BACKEND API", "ASP.NET Core 9 (C#)\n\n\u2022 ChatController\n\u2022 AnalysisController\n\u2022 AiService\n\u2022 SqlExecutionService\n\u2022 SchemaProvider", PRIMARY),
    (Inches(9.0), Inches(1.6), Inches(3.7), Inches(2.6), "AI / LLM", "GPT-4o\n(OpenAI / GitHub Models)\n\n\u2022 NL \u2192 SQL Translation\n\u2022 Result Summarization\n\u2022 Insight Generation\n\u2022 Temp: 0.1 (deterministic)", ACCENT),
]

for x, y, w, h, title, desc, color in components:
    add_shape_bg(slide, x, y, w, h, color)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.4),
                 title, font_size=16, color=WHITE, bold=True)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.5), w - Inches(0.4), h - Inches(0.6),
                 desc, font_size=11, color=RGBColor(230, 230, 255))

# Arrows between components
add_text_box(slide, Inches(3.8), Inches(2.5), Inches(0.7), Inches(0.5),
             "\u25c4\u25ba", font_size=20, color=TEXT_MED, align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(3.8), Inches(2.9), Inches(0.7), Inches(0.3),
             "REST/JSON", font_size=9, color=TEXT_MED, align=PP_ALIGN.CENTER)

add_text_box(slide, Inches(8.3), Inches(2.5), Inches(0.7), Inches(0.5),
             "\u25c4\u25ba", font_size=20, color=TEXT_MED, align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(8.3), Inches(2.9), Inches(0.7), Inches(0.3),
             "HTTP", font_size=9, color=TEXT_MED, align=PP_ALIGN.CENTER)

# Database box
db_x, db_y = Inches(4.5), Inches(4.8)
add_shape_bg(slide, db_x, db_y, Inches(3.8), Inches(2.3), DARK)
add_text_box(slide, db_x + Inches(0.2), db_y + Inches(0.12), Inches(3.4), Inches(0.4),
             "SQL SERVER \u2014 ProfileTailor DB", font_size=14, color=WHITE, bold=True)
tables_text = "\u2022 SapRoles          \u2022 Users\n\u2022 SapUserRoles    \u2022 RoleTypes\n\u2022 UsersGroups"
add_text_box(slide, db_x + Inches(0.2), db_y + Inches(0.6), Inches(3.4), Inches(1.5),
             tables_text, font_size=12, color=RGBColor(200, 200, 220))

# Arrow from backend to DB
add_text_box(slide, Inches(5.9), Inches(4.3), Inches(0.8), Inches(0.5),
             "\u25bc", font_size=22, color=TEXT_MED, align=PP_ALIGN.CENTER)
add_text_box(slide, Inches(5.5), Inches(4.55), Inches(1.6), Inches(0.3),
             "T-SQL (SELECT only)", font_size=9, color=TEXT_MED, align=PP_ALIGN.CENTER)

# Tech stack table
add_text_box(slide, Inches(9.0), Inches(4.8), Inches(3.7), Inches(0.4),
             "Technology Stack", font_size=16, color=DARK, bold=True)

stack = [
    ("Frontend", "React 18 + TS + Vite"),
    ("Backend", "ASP.NET Core 9 (C#)"),
    ("AI Engine", "GPT-4o (OpenAI)"),
    ("Database", "SQL Server"),
]
for idx, (layer, tech) in enumerate(stack):
    y = Inches(5.3) + Inches(idx * 0.42)
    add_text_box(slide, Inches(9.0), y, Inches(1.6), Inches(0.35),
                 layer, font_size=12, color=PRIMARY, bold=True)
    add_text_box(slide, Inches(10.6), y, Inches(2.2), Inches(0.35),
                 tech, font_size=12, color=TEXT_DARK)

slide_number_footer(slide, 6, 10)


# ============================================================
# SLIDE 7: Key Features
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), PRIMARY)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Key Features", font_size=36, color=DARK, bold=True)

# Feature 1: NL Chat
add_shape_bg(slide, Inches(0.6), Inches(1.4), Inches(5.9), Inches(2.4), RGBColor(238, 242, 255))
add_shape_bg(slide, Inches(0.6), Inches(1.4), Inches(0.08), Inches(2.4), PRIMARY)
add_text_box(slide, Inches(0.9), Inches(1.5), Inches(5.3), Inches(0.4),
             "\U0001f4ac  Natural Language Chat", font_size=20, color=PRIMARY, bold=True)
txBox = add_text_box(slide, Inches(0.9), Inches(1.95), Inches(5.3), Inches(1.7), "", font_size=13, color=TEXT_MED)
tf = txBox.text_frame
tf.paragraphs[0].text = "Ask in plain English \u2014 no SQL needed:"
tf.paragraphs[0].font.size = Pt(13)
tf.paragraphs[0].font.color.rgb = TEXT_MED
tf.paragraphs[0].font.name = "Segoe UI"
for q in ['"Show me all users with more than 10 roles"',
          '"Find duplicate roles in the system"',
          '"Which users in Finance have roles expiring?"',
          '"List roles that grant all transaction codes"']:
    p = add_bullet(tf, q, font_size=12, color=TEXT_DARK)

# Feature 2: 7 Pre-Built Analyses
add_shape_bg(slide, Inches(6.8), Inches(1.4), Inches(5.9), Inches(2.4), RGBColor(236, 253, 245))
add_shape_bg(slide, Inches(6.8), Inches(1.4), Inches(0.08), Inches(2.4), SUCCESS)
add_text_box(slide, Inches(7.1), Inches(1.5), Inches(5.3), Inches(0.4),
             "\U0001f4ca  7 Pre-Built Governance Analyses", font_size=20, color=SUCCESS, bold=True)
analyses = ["Duplicate Roles", "Unused Roles", "Over-Privileged Users",
            "Expired Assignments", "Orphaned Assignments", "Role Distribution", "User Risk Scoring"]
txBox2 = add_text_box(slide, Inches(7.1), Inches(1.95), Inches(5.3), Inches(1.7), "", font_size=12, color=TEXT_DARK)
tf2 = txBox2.text_frame
tf2.paragraphs[0].text = analyses[0]
tf2.paragraphs[0].font.size = Pt(12)
tf2.paragraphs[0].font.color.rgb = TEXT_DARK
tf2.paragraphs[0].font.name = "Segoe UI"
for a in analyses[1:]:
    add_bullet(tf2, a, font_size=12, color=TEXT_DARK)

# Feature 3: AI Insights with Severity
add_shape_bg(slide, Inches(0.6), Inches(4.1), Inches(3.8), Inches(3.0), RGBColor(254, 243, 199))
add_shape_bg(slide, Inches(0.6), Inches(4.1), Inches(0.08), Inches(3.0), WARNING)
add_text_box(slide, Inches(0.9), Inches(4.2), Inches(3.3), Inches(0.4),
             "\U0001f9e0  AI Insights", font_size=18, color=DARK, bold=True)
add_text_box(slide, Inches(0.9), Inches(4.65), Inches(3.3), Inches(0.35),
             "Severity-rated findings:", font_size=13, color=TEXT_MED)
add_text_box(slide, Inches(0.9), Inches(5.05), Inches(3.3), Inches(0.35),
             "\U0001f535  Info \u2014 Informational", font_size=13, color=INFO_BLUE)
add_text_box(slide, Inches(0.9), Inches(5.4), Inches(3.3), Inches(0.35),
             "\U0001f7e1  Warning \u2014 Attention needed", font_size=13, color=WARNING)
add_text_box(slide, Inches(0.9), Inches(5.75), Inches(3.3), Inches(0.35),
             "\U0001f534  Critical \u2014 Immediate action", font_size=13, color=DANGER)

# Feature 4: Safety & Security
add_shape_bg(slide, Inches(4.7), Inches(4.1), Inches(4.0), Inches(3.0), RGBColor(254, 226, 226))
add_shape_bg(slide, Inches(4.7), Inches(4.1), Inches(0.08), Inches(3.0), DANGER)
add_text_box(slide, Inches(5.0), Inches(4.2), Inches(3.5), Inches(0.4),
             "\U0001f512  Safety & Security", font_size=18, color=DARK, bold=True)
safety = ["SELECT-only enforcement", "Dangerous keyword blocking",
          "Low temperature (0.1)", "30-second query timeout", "Full schema context for AI"]
txBox3 = add_text_box(slide, Inches(5.0), Inches(4.65), Inches(3.5), Inches(2.2), "", font_size=12, color=TEXT_DARK)
tf3 = txBox3.text_frame
tf3.paragraphs[0].text = safety[0]
tf3.paragraphs[0].font.size = Pt(12)
tf3.paragraphs[0].font.color.rgb = TEXT_DARK
tf3.paragraphs[0].font.name = "Segoe UI"
for s in safety[1:]:
    add_bullet(tf3, s, font_size=12, color=TEXT_DARK)

# Feature 5: Risk Scoring
add_shape_bg(slide, Inches(9.0), Inches(4.1), Inches(3.7), Inches(3.0), RGBColor(238, 242, 255))
add_shape_bg(slide, Inches(9.0), Inches(4.1), Inches(0.08), Inches(3.0), ACCENT)
add_text_box(slide, Inches(9.3), Inches(4.2), Inches(3.2), Inches(0.4),
             "\U0001f4c8  Risk Scoring", font_size=18, color=DARK, bold=True)
add_text_box(slide, Inches(9.3), Inches(4.7), Inches(3.2), Inches(0.4),
             "Formula:", font_size=13, color=TEXT_MED)
risk_bg = add_shape_bg(slide, Inches(9.3), Inches(5.1), Inches(3.2), Inches(0.7), DARK)
add_text_box(slide, Inches(9.4), Inches(5.15), Inches(3.0), Inches(0.6),
             "Risk = Roles\u00d71\n+ Critical\u00d710\n+ Expired\u00d75",
             font_size=12, color=RGBColor(180, 255, 180), font_name="Consolas")
add_text_box(slide, Inches(9.3), Inches(5.9), Inches(3.2), Inches(0.5),
             "Enables prioritized remediation \u2014 fix highest-risk users first",
             font_size=11, color=TEXT_MED)

slide_number_footer(slide, 7, 10)


# ============================================================
# SLIDE 8: Business Impact
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), SUCCESS)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Business Impact", font_size=36, color=DARK, bold=True)

# Before/After table
data = [
    ["Metric", "Before Role Assistant", "After Role Assistant"],
    ["Time to Answer", "Days to weeks (manual SQL)", "Seconds (AI-generated)"],
    ["Audience", "DBAs & IT only", "All business users"],
    ["Analysis Coverage", "Ad-hoc, inconsistent", "7 standardized + unlimited"],
    ["Audit Preparation", "Manual role reviews", "Continuous, automated"],
    ["Risk Visibility", "Periodic (quarterly)", "Real-time, on-demand"],
]
add_table(slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(2.5), 6, 3, data,
          col_widths=[Inches(2.8), Inches(4.6), Inches(4.7)])

# Strategic benefits
add_text_box(slide, Inches(0.8), Inches(4.2), Inches(5), Inches(0.5),
             "Strategic Benefits", font_size=22, color=DARK, bold=True)

benefits = [
    ("\U0001f510", "Reduced Security Risk", "Instant detection of over-privileged users & orphaned roles", DANGER),
    ("\U0001f4cb", "Audit Readiness", "Aligns with SOX, SOD, and GRC compliance", PRIMARY),
    ("\U0001f4b0", "Cost Savings", "60-80% reduction in manual analysis effort", SUCCESS),
    ("\U0001f465", "Self-Service", "Business users get answers independently", INFO_BLUE),
    ("\u26a1", "Faster Remediation", "Risk-scored prioritization of cleanup actions", WARNING),
    ("\U0001f4ca", "Data-Driven Decisions", "Quantitative role governance", ACCENT),
]

for idx, (icon, title, desc, color) in enumerate(benefits):
    col = idx % 3
    row = idx // 3
    x = Inches(0.6) + Inches(col * 4.15)
    y = Inches(4.8) + Inches(row * 1.3)

    add_shape_bg(slide, x, y, Inches(3.9), Inches(1.1), LIGHT_GRAY)
    add_shape_bg(slide, x, y, Inches(0.06), Inches(1.1), color)

    add_text_box(slide, x + Inches(0.2), y + Inches(0.1), Inches(3.5), Inches(0.35),
                 f"{icon}  {title}", font_size=15, color=DARK, bold=True)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.5), Inches(3.5), Inches(0.5),
                 desc, font_size=12, color=TEXT_MED)

slide_number_footer(slide, 8, 10)


# ============================================================
# SLIDE 9: Future Roadmap
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.06), ACCENT)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(10), Inches(0.7),
             "Future Roadmap", font_size=36, color=DARK, bold=True)

roadmap = [
    ("v2.0", [
        ("Remediation Actions", "One-click role revocation and reassignment"),
        ("Scheduled Reports", "Automated weekly/monthly governance reports"),
    ], PRIMARY),
    ("v2.5", [
        ("Multi-System Support", "Analyze roles across multiple SAP systems"),
        ("SOD Rule Engine", "Built-in Segregation of Duties conflict detection"),
    ], ACCENT),
    ("v3.0", [
        ("AI Role Mining", "AI-suggested optimal role definitions"),
        ("Pathlock GRC Integration", "Feed insights into the Pathlock platform"),
    ], SUCCESS),
]

for idx, (version, items, color) in enumerate(roadmap):
    x = Inches(0.6) + Inches(idx * 4.15)
    y = Inches(1.5)

    # Version header
    header = add_shape_bg(slide, x, y, Inches(3.9), Inches(0.7), color)
    add_text_box(slide, x + Inches(0.2), y + Inches(0.1), Inches(3.5), Inches(0.5),
                 version, font_size=24, color=WHITE, bold=True)

    # Items
    for i, (title, desc) in enumerate(items):
        iy = y + Inches(0.9) + Inches(i * 2.2)
        card = add_shape_bg(slide, x, iy, Inches(3.9), Inches(2.0), LIGHT_GRAY)
        add_shape_bg(slide, x, iy, Inches(0.06), Inches(2.0), color)
        add_text_box(slide, x + Inches(0.2), iy + Inches(0.15), Inches(3.5), Inches(0.45),
                     title, font_size=18, color=DARK, bold=True)
        add_text_box(slide, x + Inches(0.2), iy + Inches(0.65), Inches(3.5), Inches(1.2),
                     desc, font_size=14, color=TEXT_MED)

slide_number_footer(slide, 9, 10)


# ============================================================
# SLIDE 10: Thank You / Summary
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)
add_shape_bg(slide, Inches(0), Inches(0), W, Inches(0.08), PRIMARY)

add_text_box(slide, Inches(1.5), Inches(0.8), Inches(10.3), Inches(0.8),
             "Summary", font_size=40, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

# Summary cards
summary_items = [
    ("Problem", "SAP role governance is manual, slow, and inaccessible to business users", DANGER),
    ("Solution", "AI-powered chat interface: English \u2192 SQL \u2192 Insights", PRIMARY),
    ("Technology", "React + ASP.NET Core 9 + GPT-4o + SQL Server", INFO_BLUE),
    ("Innovation", "Natural language to safe, validated SQL with governance insights", ACCENT),
    ("Impact", "Seconds instead of weeks \u00b7 Self-service \u00b7 Continuous compliance", SUCCESS),
    ("Differentiator", "Safety-first design + quantitative risk scoring", WARNING),
]

for idx, (label, text, color) in enumerate(summary_items):
    col = idx % 2
    row = idx // 2
    x = Inches(0.8) + Inches(col * 6.2)
    y = Inches(1.9) + Inches(row * 1.4)

    card = add_shape_bg(slide, x, y, Inches(5.8), Inches(1.15), RGBColor(45, 45, 65))
    add_shape_bg(slide, x, y, Inches(0.08), Inches(1.15), color)
    add_text_box(slide, x + Inches(0.25), y + Inches(0.1), Inches(5.3), Inches(0.35),
                 label, font_size=15, color=color, bold=True)
    add_text_box(slide, x + Inches(0.25), y + Inches(0.5), Inches(5.3), Inches(0.55),
                 text, font_size=14, color=RGBColor(200, 200, 220))

# Tagline
add_text_box(slide, Inches(2), Inches(6.0), Inches(9.3), Inches(0.6),
             "\U0001f6e1\ufe0f  Role Assistant \u2014 Ask your data, don\u2019t query it.",
             font_size=24, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER)

# Thank you
add_text_box(slide, Inches(4), Inches(6.7), Inches(5.3), Inches(0.5),
             "Thank You! \U0001f64f   Questions?",
             font_size=20, color=RGBColor(180, 180, 220), align=PP_ALIGN.CENTER)

add_shape_bg(slide, Inches(0), Inches(7.42), W, Inches(0.08), PRIMARY)

slide_number_footer(slide, 10, 10)

# ============================================================
# SAVE
# ============================================================
output_path = r"c:\Projects\Hackathon\RoleAssistant_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
